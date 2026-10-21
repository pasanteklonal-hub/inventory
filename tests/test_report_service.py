import json
from io import BytesIO
from app import db
from app.models import Categoria, Item, Campo
from app.services.report_service import (
    build_reporte_semanal,
    export_reporte_semanal_xlsx,
    export_reporte_semanal_pdf,
    export_inventario_json,
)
from flask import request


def test_weekly_report_details_only_show_category_data_and_stock(app):
    categoria = Categoria(nombre='Detalle Reporte')
    db.session.add(categoria)
    db.session.flush()
    item = Item(
        categoria_id=categoria.id,
        datos={'Descripción': 'Monitor', 'Modelo': 'M24'},
        stock=4,
    )
    db.session.add(item)
    db.session.commit()

    generado_el = '20/10/2026 15:30'
    with app.test_request_context('/reporte-semanal'):
        response = export_reporte_semanal_xlsx(
            [{'categoria': categoria.nombre, 'stock': 4}],
            [item],
            generado_el,
        )
        response.direct_passthrough = False

        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(response.get_data()), read_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        detail_header_index = rows.index(('Categoría', 'Datos', 'Stock'))

        assert rows[1][:2] == ('Generado el', generado_el)
        assert rows[detail_header_index + 1] == (
            categoria.nombre,
            'Descripción: Monitor / Modelo: M24',
            4,
        )


def test_report_exports(app):
    # prepare data
    c = Categoria(nombre='RepCat')
    db.session.add(c)
    db.session.flush()

    Campo(categoria_id=c.id, nombre='Name', tipo='text')
    db.session.commit()

    it = Item(categoria_id=c.id, datos={'Name': 'X'}, stock=5)
    db.session.add(it)
    db.session.commit()

    # build report within request context
    with app.test_request_context('/reporte-semanal'):
        categorias, summary, items, selected = build_reporte_semanal(request)
        assert isinstance(summary, list)
        assert next(row['stock'] for row in summary if row['categoria'] == 'RepCat') == 5

        # xlsx
        resp_xlsx = export_reporte_semanal_xlsx(summary, items)
        assert resp_xlsx.status_code == 200
        assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in resp_xlsx.mimetype
        resp_xlsx.direct_passthrough = False
        data = resp_xlsx.get_data()
        assert len(data) > 0

        # pdf
        resp_pdf = export_reporte_semanal_pdf(summary, items)
        assert resp_pdf.status_code == 200
        assert 'application/pdf' in resp_pdf.mimetype
        resp_pdf.direct_passthrough = False
        data = resp_pdf.get_data()
        assert len(data) > 0

        # export json
        resp_json = export_inventario_json()
        assert resp_json.status_code == 200
        assert 'application/json' in resp_json.mimetype
        resp_json.direct_passthrough = False
        content = resp_json.get_data()
        parsed = json.loads(content.decode('utf-8'))
        assert isinstance(parsed, dict)
        assert 'RepCat' in parsed
        assert parsed['RepCat'][0]['__inventory_stock__'] == 5

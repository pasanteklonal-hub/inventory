import json
from io import BytesIO
from datetime import datetime
from flask import send_file
from sqlalchemy import func
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, LongTable
from reportlab.lib.units import inch
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from ..models import Categoria, Campo, Item, Movimiento
from .. import db
from ..dao.item_dao import sort_items_alphabetically
from .item_service import validate_quantity


def build_reporte_semanal(request):
    categorias = Categoria.query.filter_by(deleted_at=None).order_by(
        func.lower(Categoria.nombre),
        Categoria.id,
    ).all()
    selected = request.values.getlist('categoria_id', type=int)

    if selected:
        items = Item.query.filter(
            Item.deleted_at.is_(None),
            Item.eliminado.is_(False),
            Item.categoria_id.in_(selected),
        ).all()
    else:
        items = Item.query.filter(Item.deleted_at.is_(None), Item.eliminado.is_(False)).all()
        selected = [c.id for c in categorias]
    items = sort_items_alphabetically(items)

    summary = []
    for categoria in categorias:
        if categoria.id not in selected:
            continue
        items_categoria = Item.query.filter_by(
            categoria_id=categoria.id,
            deleted_at=None,
            eliminado=False,
        ).all()
        summary.append({'categoria': categoria.nombre, 'stock': sum(item.stock for item in items_categoria)})

    return categorias, summary, items, selected


def export_reporte_semanal_xlsx(summary, items, generado_el=None):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Inventario Semanal'
    title_font = Font(size=14, bold=True)
    data_font = Font(size=14)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.append(['Reporte Semanal de Inventario'])
    ws['A1'].font = Font(size=16, bold=True)
    ws.append(['Generado el', generado_el or datetime.now().strftime('%d/%m/%Y %H:%M')])
    ws.append([])
    ws.append(['Resumen por Categoría'])
    ws['A4'].font = title_font
    ws.append(['Categoría', 'Stock Actual'])
    for cell in ws[5]:
        cell.font = title_font
        cell.alignment = center
    for row in summary:
        ws.append([row['categoria'], row['stock']])
        ws.cell(row=ws.max_row, column=1).font = data_font
        ws.cell(row=ws.max_row, column=2).font = data_font

    ws.append([])
    ws.append(['Detalle de Items'])
    ws['A' + str(ws.max_row)].font = title_font
    ws.append(['Categoría', 'Datos', 'Stock'])
    for cell in ws[ws.max_row]:
        cell.font = title_font
        cell.alignment = center
    for item in items:
        datos_texto = ' / '.join(f"{k}: {v}" for k, v in (item.datos or {}).items())
        ws.append([
            item.categoria.nombre,
            datos_texto,
            item.stock,
        ])
        for col in range(1, 4):
            ws.cell(row=ws.max_row, column=col).font = data_font
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(wrap_text=True, vertical='top')

    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column_letter
        for cell in column_cells:
            try:
                cell_value = str(cell.value) if cell.value is not None else ''
                max_length = max(max_length, len(cell_value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length * 1.2, 60)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        download_name='reporte_semanal_inventario.xlsx',
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def build_reporte_semanal_pdf(summary, items, generado_el=None):
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='ReportTitle', fontName='Helvetica-Bold', fontSize=18, leading=22, spaceAfter=12))
    styles.add(ParagraphStyle(name='ReportSubtitle', fontName='Helvetica', fontSize=14, leading=18, spaceAfter=12))
    styles.add(ParagraphStyle(name='ReportNormal', fontName='Helvetica', fontSize=14, leading=18))

    elements = []
    elements.append(Paragraph('Reporte Semanal de Inventario', styles['ReportTitle']))
    elements.append(Paragraph(
        f"Fecha de generación: {generado_el or datetime.now().strftime('%d/%m/%Y %H:%M')}",
        styles['ReportSubtitle'],
    ))
    elements.append(
        Paragraph(
            f"Total stock actual: {sum(row['stock'] for row in summary)}",
            styles['ReportNormal'],
        )
    )
    elements.append(Spacer(1, 12))

    summary_data = [['Categoría', 'Stock Actual']]
    for row in summary:
        summary_data.append([row['categoria'], row['stock']])
    summary_table = Table(summary_data, colWidths=[3.5 * inch, 2.5 * inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f81bd')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 14),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph('Detalle de Items', styles['ReportSubtitle']))
    items_data = [['Categoría', 'Datos', 'Stock']]
    for item in items:
        datos_texto = '<br/>'.join(f'{k}: {v}' for k, v in (item.datos or {}).items())
        items_data.append([
            Paragraph(item.categoria.nombre, styles['ReportNormal']),
            Paragraph(datos_texto or 'Sin datos', styles['ReportNormal']),
            Paragraph(str(item.stock), styles['ReportNormal']),
        ])

    col_widths = [1.5 * inch, 5.0 * inch, 0.7 * inch]
    items_table = LongTable(items_data, colWidths=col_widths, repeatRows=1, splitByRow=1)
    items_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f81bd')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 14),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('FONTNAME', (1, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (1, 1), (-1, -1), 14),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(items_table)

    doc.build(elements)
    output.seek(0)
    return output.getvalue()


def export_reporte_semanal_pdf(summary, items, generado_el=None):
    output = BytesIO(build_reporte_semanal_pdf(summary, items, generado_el))

    return send_file(
        output,
        download_name='reporte_semanal_inventario.pdf',
        as_attachment=True,
        mimetype='application/pdf',
    )


def export_inventario_json():
    categorias = Categoria.query.filter_by(deleted_at=None).order_by(
        func.lower(Categoria.nombre),
        Categoria.id,
    ).all()
    out = {}
    for cat in categorias:
        items = sort_items_alphabetically(
            Item.query.filter_by(categoria_id=cat.id, deleted_at=None, eliminado=False).all()
        )
        out[cat.nombre] = [
            {'categoria': cat.nombre, **(item.datos or {}), '__inventory_stock__': item.stock}
            for item in items
        ]

    bio = BytesIO()
    bio.write(json.dumps(out, ensure_ascii=False, indent=2).encode('utf-8'))
    bio.seek(0)
    return send_file(
        bio,
        download_name='export_inventario.json',
        as_attachment=True,
        mimetype='application/json',
    )


def import_inventario_json(request):
    if request.method != 'POST':
        return None

    file = request.files.get('file')
    if not file:
        from flask import flash, redirect, url_for
        flash('No se subió ningún archivo', 'danger')
        return redirect(url_for('main.importar_json'))

    filename = file.filename
    if not filename or not filename.lower().endswith('.json'):
        from flask import flash, redirect, url_for
        flash('Debe subir un archivo .json', 'danger')
        return redirect(url_for('main.importar_json'))

    try:
        data = json.loads(file.read().decode('utf-8'))
    except Exception as e:
        from flask import flash, redirect, url_for
        flash(f'Error leyendo JSON: {e}', 'danger')
        return redirect(url_for('main.importar_json'))

    try:
        stock_by_category = {
            categoria_name: [
                validate_quantity(item.get('__inventory_stock__', 1))
                for item in items
            ]
            for categoria_name, items in data.items()
        }
    except (AttributeError, TypeError, ValueError) as error:
        from flask import flash, redirect, url_for
        flash(f'Archivo JSON inválido: cantidad de stock incorrecta ({error})', 'danger')
        return redirect(url_for('main.importar_json'))

    Movimiento.query.delete()
    Item.query.delete()
    Campo.query.delete()
    Categoria.query.delete()
    db.session.commit()
    db.session.expire_all()

    for categoria_name, items in data.items():
        categoria = Categoria(nombre=categoria_name, descripcion='Importada desde JSON (UI)')
        db.session.add(categoria)
        db.session.flush()

        campos_nombres = []
        for it in items:
            for k in it.keys():
                if k.strip().lower() == 'categoria' or k == '__inventory_stock__':
                    continue
                if k not in campos_nombres:
                    campos_nombres.append(k)

        for idx, nombre in enumerate(campos_nombres):
            campo = Campo(
                categoria_id=categoria.id,
                nombre=nombre,
                tipo='text',
                requerido=False,
                orden=idx,
            )
            db.session.add(campo)

        for it, stock in zip(items, stock_by_category[categoria_name]):
            item_data = {
                k: v for k, v in it.items()
                if k.strip().lower() != 'categoria' and k != '__inventory_stock__'
            }
            item = Item(categoria_id=categoria.id, datos=item_data, stock=stock)
            db.session.add(item)

    db.session.commit()
    from flask import flash, redirect, url_for
    flash('Importación completada correctamente', 'success')
    return redirect(url_for('main.equipos'))


def list_movimientos():
    return Movimiento.query.order_by(Movimiento.created_at.desc()).all()
